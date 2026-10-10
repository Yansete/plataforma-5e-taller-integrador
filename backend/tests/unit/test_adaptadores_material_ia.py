"""Adaptadores de salida de la ingesta y de la IA, sin red: archivos reales en memoria y HTTP simulado."""
import io
import json

import httpx
import pytest

from plataforma5e.adapters.outbound.extraccion.extractor_documentos import ExtractorDocumentos
from plataforma5e.adapters.outbound.fuentes.wikipedia import FuenteWikipedia
from plataforma5e.adapters.outbound.ia.configuracion_ia import crear_generadores, describir
from plataforma5e.adapters.outbound.ia.generador_llm import ErrorProveedorIA, GeneradorLLM
from plataforma5e.adapters.outbound.ia.generador_reglas import GeneradorReglas
from plataforma5e.application.services.redaccion import construir_contexto, normalizar_recursos
from plataforma5e.domain.material import MaterialError

MATERIAL = (
    'Una red de computadoras es un conjunto de equipos conectados que comparten recursos e información. '
    'El modelo OSI es un modelo de referencia que divide la comunicación en siete capas. '
    'El modelo TCP/IP es el modelo que se usa en internet y tiene cuatro capas. '
    'El protocolo TCP es orientado a la conexión y garantiza una entrega confiable de los datos. '
    'El protocolo UDP no establece conexión y por eso es más rápido para videollamadas. '
    'El enrutador es el dispositivo que decide por qué camino viaja cada paquete hasta su destino. '
    'La encapsulación es el proceso por el cual cada capa agrega su propio encabezado a los datos.'
)


def pdf_minimo(texto: str) -> bytes:
    """PDF de una página con texto (sin tildes, fuente Helvetica)."""
    flujo = f'BT /F1 12 Tf 72 720 Td ({texto}) Tj ET'.encode()
    objetos = [b'<< /Type /Catalog /Pages 2 0 R >>', b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
               b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>',
               b'<< /Length %d >>\nstream\n' % len(flujo) + flujo + b'\nendstream',
               b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>']
    salida, posiciones = io.BytesIO(), []
    salida.write(b'%PDF-1.4\n')
    for i, objeto in enumerate(objetos, start=1):
        posiciones.append(salida.tell())
        salida.write(b'%d 0 obj\n' % i + objeto + b'\nendobj\n')
    xref = salida.tell()
    salida.write(b'xref\n0 %d\n0000000000 65535 f \n' % (len(objetos) + 1))
    salida.writelines(b'%010d 00000 n \n' % p for p in posiciones)
    salida.write(b'trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF' % (len(objetos) + 1, xref))
    return salida.getvalue()


def pptx_minimo() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches
    presentacion = Presentation()
    diapositiva = presentacion.slides.add_slide(presentacion.slide_layouts[1])
    diapositiva.shapes.title.text = 'Protocolos'
    diapositiva.placeholders[1].text = 'TCP garantiza la entrega confiable de los datos en la red.'
    diapositiva.shapes.add_table(1, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table.cell(0, 0).text = 'Capa de transporte'
    diapositiva.notes_slide.notes_text_frame.text = 'Explicar con un ejemplo.'
    presentacion.slides.add_slide(presentacion.slide_layouts[6])  # diapositiva en blanco
    salida = io.BytesIO()
    presentacion.save(salida)
    return salida.getvalue()


def test_extractor_pdf_pptx_y_errores():
    extractor = ExtractorDocumentos()
    assert extractor.extraer('pdf', pdf_minimo('El protocolo TCP garantiza la entrega.')) == [('p. 1', 'El protocolo TCP garantiza la entrega.')]
    paginas = extractor.extraer('pptx', pptx_minimo())
    assert paginas[0][0] == 'diapositiva 1' and 'TCP garantiza' in paginas[0][1]
    assert 'Capa de transporte' in paginas[0][1] and 'Notas: Explicar' in paginas[0][1] and paginas[1] == ('diapositiva 2', '')
    for tipo, contenido, codigo in [('pdf', b'%PDF-1.4 roto', 'ARCHIVO_ILEGIBLE'), ('pptx', b'no es zip', 'ARCHIVO_ILEGIBLE'),
                                    ('txt', b'\xff\xfe\x00', 'FORMATO_INVALIDO'), ('doc', b'x', 'FORMATO_INVALIDO')]:
        with pytest.raises(MaterialError) as error:
            extractor.extraer(tipo, contenido)
        assert error.value.codigo == codigo


def test_extractor_txt_por_secciones_y_por_titulos_de_wikipedia():
    extractor = ExtractorDocumentos()
    largo = '\n\n'.join(['Párrafo de prueba con varias palabras. ' * 20] * 6)
    secciones = extractor.extraer('txt', largo.encode('utf-8-sig'))
    assert len(secciones) > 1 and secciones[0][0] == 'sección 1'
    wiki = 'Introducción del artículo.\n\n== Historia ==\nOrigen de la red.\n\n=== Detalle ===\nMás datos.'
    secciones = [(u, t.strip()) for u, t in extractor.extraer('txt', wiki.encode())]
    assert secciones == [('introducción', 'Introducción del artículo.'), ('sección «Historia»', 'Origen de la red.'),
                         ('sección «Detalle»', 'Más datos.')]


def transporte(respuestas):
    """Cliente httpx que responde en orden y guarda las peticiones."""
    pedidos = []

    def manejar(peticion):
        pedidos.append(peticion)
        estado, cuerpo = respuestas[min(len(pedidos), len(respuestas)) - 1]
        return httpx.Response(estado, json=cuerpo)
    return httpx.Client(transport=httpx.MockTransport(manejar)), pedidos


CONTEXTO = construir_contexto({'name': 'Redes'}, {'id': 'u1', 'number': 1, 'title': 'Protocolos', 'outcomes': []},
                              {'stage': 'evaluate', 'resourceType': 'item_opcion_multiple', 'quantity': 1, 'optionCount': 4},
                              [{'id': 'f1', 'documento_id': 'd', 'ubicacion': 'p. 1', 'texto': MATERIAL}], {'d': 'apuntes.txt'})
RESPUESTA = {'recursos': [{'titulo': 'T', 'contenido': 'C', 'citas': [{'afirmacion': 'a', 'fragmentos': ['F1']}], 'alternativas': []}]}


def test_gemini_envia_instrucciones_y_lee_el_texto():
    cliente, pedidos = transporte([(200, {'candidates': [{'content': {'parts': [{'text': json.dumps(RESPUESTA)}]}}]})])
    salida = GeneradorLLM('gemini', 'clave', cliente=cliente).generar(CONTEXTO)
    assert salida == RESPUESTA
    assert pedidos[0].url.path.endswith('/models/gemini-3.8-flash:generateContent')
    assert pedidos[0].headers['x-goog-api-key'] == 'clave'
    cuerpo = json.loads(pedidos[0].content)
    assert cuerpo['generationConfig']['responseMimeType'] == 'application/json' and '[F1]' in cuerpo['contents'][0]['parts'][0]['text']


def test_anthropic_y_openai_compatible():
    cliente, pedidos = transporte([(200, {'content': [{'type': 'text', 'text': json.dumps(RESPUESTA)}]})])
    generador = GeneradorLLM('anthropic', 'clave', cliente=cliente)
    assert generador.generar(CONTEXTO) == RESPUESTA and generador.descripcion == 'Anthropic · claude-sonnet-5-5'
    assert pedidos[0].url.path == '/v1/messages' and pedidos[0].headers['anthropic-version'] == '2023-06-01'
    cliente, pedidos = transporte([(200, {'choices': [{'message': {'content': json.dumps(RESPUESTA['recursos'])}}]})])
    generador = GeneradorLLM('openai', 'clave', modelo='llama-3.3', url_base='https://api.groq.com/openai/v1/', cliente=cliente)
    assert generador.generar(CONTEXTO) == RESPUESTA
    assert str(pedidos[0].url) == 'https://api.groq.com/openai/v1/chat/completions'
    assert pedidos[0].headers['authorization'] == 'Bearer clave'


def test_reintenta_una_vez_si_el_json_no_es_valido():
    cliente, pedidos = transporte([(200, {'content': [{'type': 'text', 'text': 'no es json'}]}),
                                   (200, {'content': [{'type': 'text', 'text': json.dumps(RESPUESTA)}]})])
    assert GeneradorLLM('anthropic', 'k', cliente=cliente).generar(CONTEXTO) == RESPUESTA
    assert len(pedidos) == 2 and 'no era JSON válido' in json.loads(pedidos[1].content)['messages'][0]['content']


@pytest.mark.parametrize('estado,cuerpo,texto', [
    (429, {'error': {'message': 'quota'}}, 'límite de uso'),
    (400, {'error': {'message': 'API key not valid'}}, 'API key not valid'),
    (200, {'promptFeedback': {'blockReason': 'SAFETY'}}, 'SAFETY'),
])
def test_errores_del_proveedor_se_explican(estado, cuerpo, texto):
    cliente, _ = transporte([(estado, cuerpo)])
    with pytest.raises(ErrorProveedorIA) as error:
        GeneradorLLM('gemini', 'k', cliente=cliente).generar(CONTEXTO)
    assert texto in str(error.value)


def test_sin_conexion_y_proveedor_desconocido():
    def caer(peticion):
        raise httpx.ConnectError('sin red')
    with pytest.raises(ErrorProveedorIA):
        GeneradorLLM('openai', 'k', cliente=httpx.Client(transport=httpx.MockTransport(caer))).generar(CONTEXTO)
    with pytest.raises(ValueError):
        GeneradorLLM('otro', 'k')


def test_configuracion_desde_el_entorno():
    principal, respaldo = crear_generadores({})
    assert principal is respaldo and principal.usa_ia is False
    assert describir(principal) == {'descripcion': 'Generador por reglas (sin IA)', 'usaIA': False, 'proveedor': 'reglas', 'modelo': None}
    principal, respaldo = crear_generadores({'IA_API_KEY': 'k'})
    assert principal.proveedor == 'gemini' and respaldo.usa_ia is False
    principal, _ = crear_generadores({'IA_API_KEY': 'k', 'IA_PROVEEDOR': 'Anthropic', 'IA_MODELO': 'claude-haiku-5-5', 'IA_TIEMPO_MAX': '30'})
    assert describir(principal)['modelo'] == 'claude-haiku-5-5'
    assert crear_generadores({'IA_API_KEY': 'k', 'IA_PROVEEDOR': 'reglas'})[0].usa_ia is False


def test_wikipedia_busca_y_limpia_el_articulo():
    def manejar(peticion):
        if peticion.url.params.get('list') == 'search':
            return httpx.Response(200, json={'query': {'search': [{'title': 'Protocolo TCP'}, {'title': 'Desambiguación'}, {'title': 'Falta'}]}})
        titulo = peticion.url.params['titles']
        if titulo == 'Falta':
            return httpx.Response(200, json={'query': {'pages': [{'missing': True}]}})
        texto = 'Corto.' if titulo == 'Desambiguación' else MATERIAL + '\n\n== Referencias ==\nLibro.'
        return httpx.Response(200, json={'query': {'pages': [{'title': titulo, 'extract': texto, 'fullurl': 'https://es.wikipedia.org/wiki/TCP'}]}})
    fuente = FuenteWikipedia(cliente=httpx.Client(transport=httpx.MockTransport(manejar)))
    articulos = fuente.buscar('protocolo TCP', 3)
    assert [a['titulo'] for a in articulos] == ['Protocolo TCP']
    assert articulos[0]['licencia'] == 'CC BY-SA 4.0' and 'Referencias' not in articulos[0]['texto']


@pytest.mark.parametrize('etapa,tipo', [
    ('evaluate', 'item_opcion_multiple'), ('engage', 'pregunta_detonante'), ('engage', 'sondeo_diagnostico'),
    ('explore', 'guia_exploracion'), ('explore', 'caso_indagacion'), ('explain', 'explicacion'),
    ('explain', 'glosario'), ('elaborate', 'ejercicio_aplicacion'),
])
def test_generador_por_reglas_produce_recursos_validos_y_citados(etapa, tipo):
    contexto = construir_contexto({'name': 'Redes'}, {'id': 'u1', 'number': 1, 'title': 'Protocolos', 'outcomes': []},
                                  {'stage': etapa, 'resourceType': tipo, 'quantity': 2, 'optionCount': 4},
                                  [{'id': 'f1', 'documento_id': 'd', 'ubicacion': 'p. 1', 'texto': MATERIAL}], {})
    recursos = normalizar_recursos(GeneradorReglas().generar(contexto), contexto)
    assert recursos, tipo
    assert all(r['citas'][0]['fragmentIds'] == ['f1'] for r in recursos)
    if tipo == 'item_opcion_multiple':
        item = recursos[0]
        assert len(item['alternativas']) == 4 and sum(a['correcta'] for a in item['alternativas']) == 1
        assert '________' in item['contenido']
        clave = next(a['texto'] for a in item['alternativas'] if a['correcta'])
        assert clave.lower() not in item['titulo'].lower()


def test_generador_por_reglas_sin_material_util_no_inventa():
    contexto = construir_contexto({}, {'id': 'u1', 'number': 1, 'title': 'U', 'outcomes': []},
                                  {'stage': 'evaluate', 'resourceType': 'item_opcion_multiple', 'quantity': 1, 'optionCount': 4},
                                  [{'id': 'f1', 'documento_id': 'd', 'ubicacion': 'p. 1', 'texto': 'Hola.'}], {})
    assert GeneradorReglas().generar(contexto) == {'recursos': []}

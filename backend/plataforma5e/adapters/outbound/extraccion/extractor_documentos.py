"""Lee el texto de PDF, PPTX y TXT (EN-007). Devuelve una entrada por página, diapositiva o sección."""
import io
import re

from plataforma5e.domain.material import MaterialError

MAX_PAGINAS = 400            # un libro largo se procesa hasta aquí
CARACTERES_POR_SECCION = 3000
_TITULO_WIKI = re.compile(r'^\s*(={2,})\s*(.+?)\s*\1\s*$', re.MULTILINE)


class ExtractorDocumentos:
    def extraer(self, tipo: str, contenido: bytes) -> list[tuple[str, str]]:
        if tipo == 'pdf':
            return self._pdf(contenido)
        if tipo == 'pptx':
            return self._pptx(contenido)
        if tipo == 'txt':
            return self._txt(contenido)
        raise MaterialError('FORMATO_INVALIDO', 'Usa PDF, PPTX o TXT.')

    @staticmethod
    def _pdf(contenido: bytes) -> list[tuple[str, str]]:
        from pypdf import PdfReader
        try:
            lector = PdfReader(io.BytesIO(contenido))
            if lector.is_encrypted and not lector.decrypt(''):
                raise MaterialError('PDF_PROTEGIDO', 'El PDF está protegido con contraseña. Quita la protección y vuelve a subirlo.')
            return [(f'p. {i}', pagina.extract_text() or '') for i, pagina in enumerate(lector.pages[:MAX_PAGINAS], start=1)]
        except MaterialError:
            raise
        except Exception as exc:  # pypdf lanza errores distintos según cómo esté dañado el archivo
            raise MaterialError('ARCHIVO_ILEGIBLE', 'No se pudo leer el PDF. Puede estar dañado; ábrelo, guárdalo de nuevo como PDF y vuelve a subirlo.') from exc

    @staticmethod
    def _pptx(contenido: bytes) -> list[tuple[str, str]]:
        from pptx import Presentation
        try:
            presentacion = Presentation(io.BytesIO(contenido))
        except Exception as exc:  # zip dañado, paquete incompleto u otro formato
            raise MaterialError('ARCHIVO_ILEGIBLE', 'No se pudo leer la presentación. Guárdala de nuevo como PPTX y vuelve a subirla.') from exc
        paginas = []
        for numero, diapositiva in enumerate(presentacion.slides, start=1):
            if numero > MAX_PAGINAS:
                break
            textos = [t for forma in diapositiva.shapes for t in _textos_de_forma(forma)]
            if diapositiva.has_notes_slide:
                notas = diapositiva.notes_slide.notes_text_frame.text.strip()
                if notas:
                    textos.append(f'Notas: {notas}')
            # Cada caja de texto suele ser una idea: se separan como párrafos.
            paginas.append((f'diapositiva {numero}', '\n\n'.join(t.strip() for t in textos if t.strip())))
        return paginas

    @staticmethod
    def _txt(contenido: bytes) -> list[tuple[str, str]]:
        try:
            texto = contenido.decode('utf-8-sig')
        except UnicodeDecodeError as exc:
            raise MaterialError('FORMATO_INVALIDO', 'Guarda el TXT con codificación UTF-8.') from exc
        titulos = list(_TITULO_WIKI.finditer(texto))
        if titulos:  # artículos de Wikipedia: cada sección conserva su título como ubicación
            secciones = [('introducción', texto[:titulos[0].start()])]
            for actual, siguiente in zip(titulos, [*titulos[1:], None]):
                fin = siguiente.start() if siguiente else len(texto)
                secciones.append((f'sección «{actual.group(2)}»', texto[actual.end():fin]))
            return [(u, t) for u, t in secciones if t.strip()][:MAX_PAGINAS]
        secciones, actual = [], ''
        for parrafo in re.split(r'\n\s*\n', texto):
            if actual and len(actual) + len(parrafo) > CARACTERES_POR_SECCION:
                secciones.append(actual)
                actual = ''
            actual = f'{actual}\n\n{parrafo}' if actual else parrafo
        if actual.strip():
            secciones.append(actual)
        return [(f'sección {i}', s) for i, s in enumerate(secciones[:MAX_PAGINAS], start=1)]


def _textos_de_forma(forma) -> list[str]:
    """Texto de cuadros, tablas y grupos de una diapositiva."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    if getattr(forma, 'shape_type', None) == MSO_SHAPE_TYPE.GROUP:
        return [t for interna in forma.shapes for t in _textos_de_forma(interna)]
    if getattr(forma, 'has_table', False) and forma.has_table:
        return [' | '.join(celda.text for celda in fila.cells) for fila in forma.table.rows]
    if getattr(forma, 'has_text_frame', False) and forma.has_text_frame:
        return [forma.text_frame.text]
    return []
